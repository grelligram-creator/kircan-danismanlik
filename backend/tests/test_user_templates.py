"""Backend tests — v6 custom Word template management (/api/user_templates/*, chat image, docx download)."""
import io
import os
import zipfile
from pathlib import Path

import pytest
import requests
from dotenv import dotenv_values

frontend_env = dotenv_values("/app/frontend/.env")
base_url = os.environ.get("REACT_APP_BACKEND_URL") or frontend_env.get("REACT_APP_BACKEND_URL")
if not base_url:
    raise RuntimeError("REACT_APP_BACKEND_URL missing")
BASE_URL = base_url.rstrip("/")

SESSION_TOKEN = "test_session_seed_001"
USER_ID = "test-user-seed-001"


@pytest.fixture(scope="module")
def client():
    s = requests.Session()
    s.headers.update({"Authorization": f"Bearer {SESSION_TOKEN}"})
    return s


@pytest.fixture(scope="module")
def seeded_template(client):
    """Reuse the already-uploaded 'Test Konut Şablonu' to conserve LLM budget."""
    r = client.get(f"{BASE_URL}/api/user_templates")
    assert r.status_code == 200, r.text
    items = r.json()["templates"]
    assert items, "No seeded user template found"
    return items[0]


# ---------------- Listing / retrieval ----------------
class TestUserTemplateRead:
    def test_list(self, client):
        r = client.get(f"{BASE_URL}/api/user_templates")
        assert r.status_code == 200, r.text
        data = r.json()
        assert isinstance(data["templates"], list) and len(data["templates"]) >= 1
        t = data["templates"][0]
        assert "original_docx_path" not in t and "prepared_docx_path" not in t
        assert "_id" not in t
        assert t["template_id"].startswith("utpl_")
        assert isinstance(t["fields"], list) and len(t["fields"]) > 0
        for f in t["fields"]:
            assert set(["key", "label", "type", "node_id"]).issubset(f.keys())
            assert f["type"] in ("text", "number", "date", "textarea", "image")

    def test_get_single(self, client, seeded_template):
        tid = seeded_template["template_id"]
        r = client.get(f"{BASE_URL}/api/user_templates/{tid}")
        assert r.status_code == 200, r.text
        t = r.json()
        assert t["template_id"] == tid
        assert t["name"] == seeded_template["name"]
        assert len(t["fields"]) == len(seeded_template["fields"])
        assert "original_docx_path" not in t and "prepared_docx_path" not in t

    def test_get_404(self, client):
        r = client.get(f"{BASE_URL}/api/user_templates/utpl_doesnotexist")
        assert r.status_code == 404, r.text

    def test_requires_auth(self):
        r = requests.get(f"{BASE_URL}/api/user_templates")
        assert r.status_code in (401, 403), r.status_code

    def test_templates_endpoint_has_custom(self, client):
        r = client.get(f"{BASE_URL}/api/templates")
        assert r.status_code == 200, r.text
        data = r.json()
        assert isinstance(data.get("templates"), list) and len(data["templates"]) >= 4
        assert isinstance(data.get("custom_templates"), list)
        assert len(data["custom_templates"]) >= 1
        assert all("_id" not in c for c in data["custom_templates"])


# ---------------- PATCH persistence ----------------
class TestUserTemplatePatch:
    def test_patch_label_and_type_persists(self, client, seeded_template):
        tid = seeded_template["template_id"]
        orig = client.get(f"{BASE_URL}/api/user_templates/{tid}").json()
        fields = [dict(f) for f in orig["fields"]]
        original_label = fields[0]["label"]
        original_type = fields[0]["type"]
        fields[0]["label"] = "TEST_Guncel Etiket"
        fields[0]["type"] = "textarea"

        r = client.patch(f"{BASE_URL}/api/user_templates/{tid}",
                         json={"fields": fields, "description": "TEST_desc"})
        assert r.status_code == 200, r.text
        upd = r.json()
        assert upd["fields"][0]["label"] == "TEST_Guncel Etiket"
        assert upd["fields"][0]["type"] == "textarea"
        assert upd["description"] == "TEST_desc"

        got = client.get(f"{BASE_URL}/api/user_templates/{tid}").json()
        assert got["fields"][0]["label"] == "TEST_Guncel Etiket"
        assert got["fields"][0]["type"] == "textarea"
        assert len(got["fields"]) == len(orig["fields"])

        # restore
        fields[0]["label"] = original_label
        fields[0]["type"] = original_type
        rb = client.patch(f"{BASE_URL}/api/user_templates/{tid}",
                          json={"fields": fields, "description": orig.get("description", "")})
        assert rb.status_code == 200, rb.text
        assert rb.json()["fields"][0]["label"] == original_label

    def test_patch_invalid_type_falls_back_to_text(self, client, seeded_template):
        tid = seeded_template["template_id"]
        orig = client.get(f"{BASE_URL}/api/user_templates/{tid}").json()
        fields = [dict(f) for f in orig["fields"]]
        bad = dict(fields[-1])
        bad["type"] = "banana"
        fields[-1] = bad
        r = client.patch(f"{BASE_URL}/api/user_templates/{tid}", json={"fields": fields})
        assert r.status_code == 200, r.text
        assert r.json()["fields"][-1]["type"] == "text"

    def test_patch_404(self, client):
        r = client.patch(f"{BASE_URL}/api/user_templates/utpl_nope", json={"name": "x"})
        assert r.status_code == 404, r.text


# ---------------- Chat creation + preview + image + download ----------------
class TestUserTemplateChatFlow:
    def test_full_flow(self, client, seeded_template):
        tid = seeded_template["template_id"]
        fields = seeded_template["fields"]
        n = len(fields)

        # 1. create chat
        r = client.post(f"{BASE_URL}/api/chats", json={"mode": "report", "user_template_id": tid})
        assert r.status_code == 200, r.text
        data = r.json()
        chat = data.get("chat", data)
        chat_id = chat["chat_id"]
        assert chat["user_template_id"] == tid
        assert chat["template_name"] == seeded_template["name"]

        msgs_greeting = data.get("messages") or []
        greeting = ""
        if msgs_greeting:
            greeting = msgs_greeting[0].get("content", "")
        else:
            mr = client.get(f"{BASE_URL}/api/chats/{chat_id}")
            assert mr.status_code == 200, mr.text
            body = mr.json()
            ms = body.get("messages", body)
            greeting = ms[0]["content"] if ms else ""
        assert seeded_template["name"] in greeting, greeting[:300]
        assert f"{n} alan" in greeting, greeting[:300]

        # 2. preview with empty values -> ph-empty
        pr = client.get(f"{BASE_URL}/api/user_templates/{tid}/preview", params={"chat_id": chat_id})
        assert pr.status_code == 200, pr.text
        pdata = pr.json()
        assert "html" in pdata and "fields" in pdata
        assert "ph-empty" in pdata["html"], pdata["html"][:500]
        assert len(pdata["fields"]) == n

        # 3. set text values directly then verify ph-filled
        text_fields = [f for f in fields if f["type"] != "image"]
        image_fields = [f for f in fields if f["type"] == "image"]
        assert text_fields, "expected at least one non-image field"
        assert image_fields, "expected at least one image field (type coverage)"

        import subprocess, json as _json
        vals = {f["key"]: f"TEST_{f['key']}_VAL" for f in text_fields[:3]}
        js = (f"db.getSiblingDB('test_database').chats.updateOne({{chat_id:'{chat_id}'}},"
              f"{{$set:{{fields:{_json.dumps(vals)}}}}})")
        subprocess.run(["mongosh", "--quiet", "--eval", js], capture_output=True, text=True)

        pr2 = client.get(f"{BASE_URL}/api/user_templates/{tid}/preview", params={"chat_id": chat_id})
        assert pr2.status_code == 200, pr2.text
        html2 = pr2.json()["html"]
        assert "ph-filled" in html2, html2[:800]
        assert list(vals.values())[0] in html2, html2[:800]

        # 4. upload image to image field
        png = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06"
               b"\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00"
               b"\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82")
        img_key = image_fields[0]["key"]
        ir = client.post(
            f"{BASE_URL}/api/chats/{chat_id}/image",
            data={"field_key": img_key},
            files={"file": ("TEST_pic.png", io.BytesIO(png), "image/png")},
        )
        assert ir.status_code == 200, ir.text
        idata = ir.json()
        assert idata["success"] is True
        assert idata["field_key"] == img_key
        assert idata["preview_url"].startswith("/api/uploads/file/")

        pr3 = client.get(f"{BASE_URL}/api/user_templates/{tid}/preview", params={"chat_id": chat_id})
        html3 = pr3.json()["html"]
        assert f'<img data-key="{img_key}"' in html3, html3[:1500]

        # 5. download docx
        dr = client.get(f"{BASE_URL}/api/chats/{chat_id}/download/docx")
        assert dr.status_code == 200, dr.text[:500]
        assert dr.content[:2] == b"PK", dr.content[:20]
        zf = zipfile.ZipFile(io.BytesIO(dr.content))
        names = zf.namelist()
        assert "word/document.xml" in names
        xml = zf.read("word/document.xml").decode("utf-8", "ignore")
        # the image field was filled, so media should exist
        assert any(nm.startswith("word/media/") for nm in names), names
        # filled text values must appear in the rendered document
        missing = [v for v in vals.values() if v not in xml]
        assert not missing, f"values missing from rendered docx: {missing}"
        assert "{{" not in xml, "unrendered Jinja token left in output docx"

        # 6. pdf not allowed for custom templates
        pdfr = client.get(f"{BASE_URL}/api/chats/{chat_id}/download/pdf")
        assert pdfr.status_code == 400, pdfr.status_code

        # 7. image upload to non-custom chat rejected
        br = client.post(f"{BASE_URL}/api/chats", json={"mode": "faq"})
        faq_id = (br.json().get("chat") or br.json())["chat_id"]
        rr = client.post(
            f"{BASE_URL}/api/chats/{faq_id}/image",
            data={"field_key": "x"},
            files={"file": ("TEST_pic.png", io.BytesIO(png), "image/png")},
        )
        assert rr.status_code == 400, rr.status_code

        # 8. non-image content type rejected
        r2 = client.post(
            f"{BASE_URL}/api/chats/{chat_id}/image",
            data={"field_key": img_key},
            files={"file": ("TEST.txt", io.BytesIO(b"hello"), "text/plain")},
        )
        assert r2.status_code == 400, r2.status_code

        # cleanup chats
        client.delete(f"{BASE_URL}/api/chats/{chat_id}")
        client.delete(f"{BASE_URL}/api/chats/{faq_id}")


# ---------------- Upload + delete (1 LLM call) ----------------
def _build_docx(path: str):
    from docx import Document
    d = Document()
    d.add_heading("TEST Konut Değerleme Formu", 0)
    d.add_paragraph("Rapor Tarihi: ____________")
    d.add_paragraph("Taşınmaz Sahibi: [YAZINIZ]")
    d.add_paragraph("Ada / Parsel: ____")
    d.add_paragraph("Genel Görünüm Fotoğrafı:")
    d.add_paragraph("")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Alan (m2)"
    t.cell(0, 1).text = ""
    t.cell(1, 0).text = "Değer (TL)"
    t.cell(1, 1).text = ""
    d.save(path)


class TestUploadAndDelete:
    def test_upload_detect_and_delete(self, client, tmp_path):
        p = tmp_path / "TEST_upload.docx"
        _build_docx(str(p))

        with open(p, "rb") as fh:
            r = client.post(
                f"{BASE_URL}/api/user_templates/upload",
                data={"name": "TEST_Upload Sablonu", "description": "TEST"},
                files={"file": ("TEST_upload.docx", fh,
                                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                timeout=180,
            )
        assert r.status_code == 200, r.text[:1000]
        body = r.json()
        tpl = body.get("template", body)
        tid = tpl.get("template_id")
        assert tid and tid.startswith("utpl_"), tpl
        assert tpl["name"] == "TEST_Upload Sablonu"
        assert "original_docx_path" not in tpl and "prepared_docx_path" not in tpl
        flds = tpl["fields"]
        assert isinstance(flds, list) and len(flds) >= 3, flds
        types = {f["type"] for f in flds}
        assert types <= {"text", "number", "date", "textarea", "image"}, types
        for f in flds:
            assert f["key"] and f["label"]
        # detection summary should have applied most fields
        summary = tpl.get("detection_summary", {})
        assert len(summary.get("applied", [])) >= 1, summary

        # preview works on the new template
        pr = client.get(f"{BASE_URL}/api/user_templates/{tid}/preview")
        assert pr.status_code == 200, pr.text
        assert "ph-empty" in pr.json()["html"], pr.json()["html"][:600]

        # non-docx upload rejected
        bad = client.post(
            f"{BASE_URL}/api/user_templates/upload",
            data={"name": "TEST_bad"},
            files={"file": ("TEST.txt", io.BytesIO(b"nope"), "text/plain")},
        )
        assert bad.status_code == 400, bad.status_code

        # delete
        dr = client.delete(f"{BASE_URL}/api/user_templates/{tid}")
        assert dr.status_code == 200, dr.text
        assert dr.json()["success"] is True
        assert client.get(f"{BASE_URL}/api/user_templates/{tid}").status_code == 404
