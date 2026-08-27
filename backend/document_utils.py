"""Document parsing and report generation utilities."""
import io
from pathlib import Path
from typing import Dict, Any, List
from datetime import datetime

from pypdf import PdfReader
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


def parse_uploaded_file(file_path: str) -> str:
    """Extract textual content from PDF/DOCX/XLSX files."""
    path = Path(file_path)
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        if suffix in (".docx", ".doc"):
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)
        if suffix in (".xlsx", ".xls"):
            wb = load_workbook(str(path), data_only=True)
            out: List[str] = []
            for sheet in wb.worksheets:
                out.append(f"# Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join(str(c) if c is not None else "" for c in row)
                    out.append(row_text)
            return "\n".join(out)
        if suffix in (".txt", ".csv", ".md"):
            return path.read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        return f"[Dosya okunamadı: {e}]"
    return "[Desteklenmeyen dosya formatı]"


def generate_pdf(report: Dict[str, Any], output_path: str) -> None:
    """Generate a PDF report from structured data."""
    doc = SimpleDocTemplate(output_path, pagesize=A4,
                            leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('title', parent=styles['Title'],
                                 fontSize=20, textColor=colors.HexColor('#0b2340'),
                                 spaceAfter=6)
    subtitle_style = ParagraphStyle('subtitle', parent=styles['Normal'],
                                    fontSize=9, textColor=colors.HexColor('#c9a24a'),
                                    spaceAfter=20)
    h2 = ParagraphStyle('h2', parent=styles['Heading2'],
                        fontSize=13, textColor=colors.HexColor('#0b2340'),
                        spaceBefore=14, spaceAfter=8)
    body = ParagraphStyle('body', parent=styles['BodyText'],
                          fontSize=10, leading=15,
                          textColor=colors.HexColor('#1a1a1a'))

    story = []
    story.append(Paragraph(report.get("template_name", "Değerleme Raporu"), title_style))
    story.append(Paragraph(
        f"KırCan Danışmanlık, Eğitim ve Değerleme Ltd. Şti. &nbsp;·&nbsp; "
        f"Rapor Tarihi: {datetime.now().strftime('%d.%m.%Y')} &nbsp;|&nbsp; Rapor No: {report.get('report_no', '-')}",
        subtitle_style))

    # Field table
    fields = report.get("fields", {})
    if fields:
        story.append(Paragraph("Gayrimenkul Bilgileri", h2))
        data = [["Alan", "Değer"]]
        for k, v in fields.items():
            data.append([str(k), str(v) if v else "-"])
        t = Table(data, colWidths=[6*cm, 10*cm])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a0a0a')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('TOPPADDING', (0, 0), (-1, 0), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e5e5')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#fafafa')]),
        ]))
        story.append(t)
        story.append(Spacer(1, 12))

    # Sections (AI generated content)
    sections = report.get("sections", {})
    for section_name, section_text in sections.items():
        story.append(Paragraph(section_name, h2))
        # split paragraphs
        for para in (section_text or "").split("\n"):
            if para.strip():
                story.append(Paragraph(para.replace("<", "&lt;").replace(">", "&gt;"), body))
                story.append(Spacer(1, 4))

    doc.build(story)


def generate_docx(report: Dict[str, Any], output_path: str) -> None:
    """Generate a DOCX report from structured data."""
    doc = Document()
    title = doc.add_heading(report.get("template_name", "Değerleme Raporu"), level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x0a, 0x0a, 0x0a)

    meta = doc.add_paragraph()
    meta_run = meta.add_run(
        f"Rapor Tarihi: {datetime.now().strftime('%d.%m.%Y')}    |    Rapor No: {report.get('report_no', '-')}"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    fields = report.get("fields", {})
    if fields:
        doc.add_heading("Gayrimenkul Bilgileri", level=2)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Alan"
        hdr[1].text = "Değer"
        for k, v in fields.items():
            row = table.add_row().cells
            row[0].text = str(k)
            row[1].text = str(v) if v else "-"

    sections = report.get("sections", {})
    for section_name, section_text in sections.items():
        doc.add_heading(section_name, level=2)
        for para in (section_text or "").split("\n"):
            if para.strip():
                doc.add_paragraph(para)

    doc.save(output_path)


def generate_udf(report: Dict[str, Any], output_path: str) -> None:
    """Generate a UYAP UDF file for court submission.

    UDF (UYAP Doküman Formatı) is a ZIP archive containing content.xml — a UYAP-specific
    XML dialect used by the Turkish courts' UYAP electronic filing system.
    """
    import zipfile
    from xml.sax.saxutils import escape as xml_escape

    template_name = report.get("template_name", "Değerleme Raporu")
    report_no = report.get("report_no", "-")

    # Build plain text body (used as CDATA content) + paragraph offsets (used as element index)
    body_lines: List[str] = []
    body_lines.append(template_name)
    body_lines.append("")
    body_lines.append(f"Rapor No: {report_no}      Tarih: {datetime.now().strftime('%d.%m.%Y')}")
    body_lines.append("")

    fields = report.get("fields", {})
    if fields:
        body_lines.append("GAYRİMENKUL BİLGİLERİ")
        for k, v in fields.items():
            body_lines.append(f"{k}: {v if v else '-'}")
        body_lines.append("")

    sections = report.get("sections", {})
    for section_name, section_text in sections.items():
        body_lines.append(section_name.upper())
        for para in (section_text or "").split("\n"):
            if para.strip():
                body_lines.append(para.strip())
        body_lines.append("")

    body_text = "\n".join(body_lines)
    # Sanitize CDATA — a ']]>' in user content would prematurely close our CDATA block
    body_text_cdata = body_text.replace("]]>", "]]]]><![CDATA[>")

    # Build elements with offsets — UYAP UDF requires paragraph elements referencing char offsets
    elements_xml_parts: List[str] = []
    offset = 0
    total_lines = len(body_lines)
    for idx, line in enumerate(body_lines):
        # All lines except the last are terminated by "\n" in the joined string
        length = len(line) + (1 if idx < total_lines - 1 else 0)
        escaped = xml_escape(line)
        # Titles (bold heuristic): all-caps lines and template_name
        is_bold = line.isupper() and line.strip() != "" or line == template_name
        family = "Times New Roman"
        size = 12 if line == template_name else 10
        elements_xml_parts.append(
            f'<paragraph Alignment="0" LineSpacing="1.15" SpaceAbove="0" SpaceBelow="0">'
            f'<content startOffset="{offset}" length="{length}" family="{family}" size="{size}" '
            f'bold="{str(is_bold).lower()}" italic="false" underline="false" '
            f'strikethrough="false" foreground="-16777216"/>'
            f'</paragraph>'
        )
        offset += length

    content_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<template format_id="1.8">\n'
        f'<content><![CDATA[{body_text_cdata}]]></content>\n'
        '<properties>\n'
        '  <pageFormat mediaSizeName="1" leftMargin="42.55" rightMargin="42.55" '
        'topMargin="42.55" bottomMargin="42.55" paperOrientation="1" headerFOffset="20.0" footerFOffset="20.0"/>\n'
        '</properties>\n'
        '<elements resolver="hvl-oluster">\n'
        + "\n".join(elements_xml_parts) + "\n"
        '</elements>\n'
        '<styles>\n'
        '  <style name="default" description="Geçerli" family="Times New Roman" size="10" '
        'bold="false" italic="false" foreground="-16777216" bgColor="-1"/>\n'
        '</styles>\n'
        '</template>\n'
    )

    with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("content.xml", content_xml)
